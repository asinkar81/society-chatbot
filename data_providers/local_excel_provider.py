"""
Local Excel-based implementation of DataProvider
"""
import logging
import pandas as pd
from pathlib import Path
from typing import List, Dict, Any, Optional
from datetime import datetime
import openpyxl
from openpyxl.utils import get_column_letter

from data_providers.base_provider import DataProvider
import config


def _ensure_columns(df, entry_columns):
    for col in entry_columns:
        if col not in df.columns:
            df[col] = None
    return df


class LocalExcelDataProvider(DataProvider):
    """Local Excel file-based data storage using openpyxl and pandas"""

    LEDGER_HEADERS = [
        "Member_ID", "Plot_No", "Date", "Particulars", "Vch_Type",
        "Vch_No", "Debit", "Credit", "Description", "Transaction_ID",
        "Transaction_Type",
    ]

    PAYMENT_REF_HEADERS = [
        "Member_ID", "Identifier_Type", "Identifier_Value",
        "Last_Seen_Date", "Confidence", "Transaction_Type",
    ]

    SUSPENSE_HEADERS = [
        "ID", "Date", "Particulars", "Amount", "Transaction_ID",
        "Transaction_Type", "Description", "Entry_Date", "Status",
        "Tagged_To", "Tagged_Date",
    ]

    EXPENSES_HEADERS = [
        "ID", "Date", "Particulars", "Amount", "Category",
        "Description", "Transaction_ID", "Entry_Date", "Member_ID",
    ]

    EXPENSE_CATEGORIES_HEADERS = [
        "ID", "Name", "Parent",
    ]

    INTEREST_INCOME_HEADERS = [
        "ID", "Date", "Particulars", "Amount", "Transaction_ID",
        "Transaction_Type", "Source_File", "Accounts_Entry_ID", "Entry_Date",
    ]

    SPLIT_RULES_HEADERS = [
        "ID", "Source_Member_ID", "Members",
    ]

    ACCOUNTS_HEADERS = [
        "Entry_ID", "Date", "Particulars", "Withdrawal", "Deposit",
        "Balance", "Calculated_Balance", "Balance_Check",
        "Transaction_Type", "Transaction_ID", "Status",
        "Source_File", "Statement_Seq",
    ]

    def __init__(self, excel_file: Path = None):
        self.excel_file = excel_file or config.SOCIETY_DATA_FILE
        self.members_sheet = "Members"
        self.ledger_sheet = "Ledger"
        self.settings_sheet = "Settings"
        self.payment_refs_sheet = "Payment_References"
        self.suspense_sheet = "Suspense_Entries"
        self.expenses_sheet = "Expenses"
        self.expense_categories_sheet = "Expense_Categories"
        self.split_rules_sheet = "Split_Rules"
        self.accounts_sheet = "Accounts"
        self.interest_income_sheet = "Interest_Income"
        self._ensure_workbook_exists()
        self._ensure_payment_refs_sheet_exists()
        self._ensure_suspense_sheet_exists()
        self._ensure_sheet(self.expenses_sheet, self.EXPENSES_HEADERS)
        self._ensure_sheet(self.expense_categories_sheet, self.EXPENSE_CATEGORIES_HEADERS)
        self._ensure_sheet(self.split_rules_sheet, self.SPLIT_RULES_HEADERS)
        self._ensure_sheet(self.accounts_sheet, self.ACCOUNTS_HEADERS)
        self._ensure_sheet(self.interest_income_sheet, self.INTEREST_INCOME_HEADERS)
        self._seed_default_categories()
        self._migrate_members_schema()
        self._migrate_expenses_schema()
        self._migrate_single_ledger()
        self._remove_current_outstanding_column()
        self._ensure_reports_sheet()
        self._ensure_processed_statements_sheet()
        self._migrate_payment_refs_schema()
        self._clean_payment_refs()
        self._cache = None

    # ── Cache ──────────────────────────────────────────────────────────────

    def _load_cache(self):
        if self._cache is not None:
            return
        try:
            mdf = pd.read_excel(self.excel_file, sheet_name=self.members_sheet)
            ldf = pd.read_excel(self.excel_file, sheet_name=self.ledger_sheet)
        except Exception:
            logging.exception("Failed to load cache")
            self._cache = {"members_list": [], "members_by_id": {}, "ledger_by_mid": {}}
            return
        cache = {}
        cache["members_list"] = mdf.to_dict("records")
        cache["members_by_id"] = {m["ID"]: m for m in cache["members_list"]}
        for m in cache["members_list"]:
            m.setdefault("WhatsApp_No", "")
        ledger_by_mid = {}
        for _, row in ldf.iterrows():
            plot_no = row.get("Plot_No")
            if pd.isna(plot_no):
                continue
            mid = int(plot_no)
            if mid not in ledger_by_mid:
                ledger_by_mid[mid] = []
            entry = row.to_dict()
            for k, v in entry.items():
                if isinstance(v, float) and v != v:
                    entry[k] = 0.0
            ledger_by_mid[mid].append(entry)
        cache["ledger_by_mid"] = ledger_by_mid
        self._cache = cache

    def _invalidate_cache(self):
        self._cache = None

    # ── Schema migrations ─────────────────────────────────────────────────

    def _migrate_members_schema(self):
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.members_sheet)
            changed = False
            if "WhatsApp_No" not in df.columns:
                df["WhatsApp_No"] = ""
                changed = True
            if changed:
                self._write_sheet(df, self.members_sheet)
        except Exception as e:
            logging.warning(f"Member schema migration: {e}")

    def _migrate_expenses_schema(self):
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.expenses_sheet)
            changed = False
            if "Transaction_Type" not in df.columns:
                df["Transaction_Type"] = ""
                changed = True
            if changed:
                self._write_sheet(df, self.expenses_sheet)
        except Exception as e:
            logging.warning(f"Expenses schema migration: {e}")

    def _backup_file(self, label: str = ""):
        import shutil
        from datetime import datetime
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        suffix = f"_{label}" if label else ""
        backup_path = config.BACKUPS_DIR / f"society_data_{ts}{suffix}.xlsx"
        shutil.copy2(self.excel_file, backup_path)
        logging.info(f"Backed up to {backup_path}")

    def _migrate_single_ledger(self):
        import re
        wb = openpyxl.load_workbook(self.excel_file)
        if "Ledger" in wb.sheetnames:
            ws = wb["Ledger"]
            headers = [cell.value for cell in ws[1]]
            if "Plot_No" in headers:
                wb.close()
                return
        wb.close()
        self._backup_file("pre_migration_single_ledger")
        master_df = pd.read_excel(self.excel_file, sheet_name=self.ledger_sheet)
        if "Plot_No" not in master_df.columns:
            master_df["Plot_No"] = None
        wb = openpyxl.load_workbook(self.excel_file)
        for sheet_name in list(wb.sheetnames):
            if re.match(r"^Plot_No_\d+$", sheet_name):
                plot_num = int(sheet_name.split("_")[-1])
                pdf = pd.read_excel(self.excel_file, sheet_name=sheet_name)
                if pdf.empty:
                    continue
                pdf["Plot_No"] = plot_num
                for _, row in pdf.iterrows():
                    match = (
                        (master_df["Date"] == row.get("Date"))
                        & (master_df["Particulars"] == row.get("Particulars"))
                        & (master_df["Vch_No"] == row.get("Vch_No"))
                        & (master_df["Credit"].fillna(0) == float(row.get("Credit", 0) or 0))
                        & (master_df["Debit"].fillna(0) == float(row.get("Debit", 0) or 0))
                    )
                    if not match.any():
                        master_df = pd.concat([master_df, pd.DataFrame([row])], ignore_index=True)
        for sheet_name in list(wb.sheetnames):
            if re.match(r"^Plot_No_\d+$", sheet_name):
                del wb[sheet_name]
        wb.close()
        self._write_sheet(master_df, self.ledger_sheet)

    def _remove_current_outstanding_column(self):
        wb = openpyxl.load_workbook(self.excel_file)
        if "Members" not in wb.sheetnames:
            wb.close()
            return
        ws = wb["Members"]
        headers = [cell.value for cell in ws[1]]
        if "Current_Outstanding" not in headers:
            wb.close()
            return
        col_idx = headers.index("Current_Outstanding") + 1
        ws.delete_cols(col_idx)
        wb.save(self.excel_file)
        wb.close()

    def _ensure_reports_sheet(self):
        wb = openpyxl.load_workbook(self.excel_file)
        if "Reports" in wb.sheetnames:
            wb.close()
            return
        ws = wb.create_sheet("Reports")
        ws.cell(row=1, column=1, value="Plot_No")
        ws.cell(row=1, column=2, value="Owner_Name")
        ws.cell(row=1, column=3, value="Current_Outstanding")
        ws.cell(row=1, column=4, value="Generated_At")
        wb.save(self.excel_file)
        wb.close()

    def _ensure_processed_statements_sheet(self):
        wb = openpyxl.load_workbook(self.excel_file)
        if "Processed_Statements" in wb.sheetnames:
            wb.close()
            return
        ws = wb.create_sheet("Processed_Statements")
        for col_idx, col_name in enumerate(
            ["Filename", "Date_Range", "Processed_At", "Entry_Count", "Format", "File_Hash"], 1
        ):
            ws.cell(row=1, column=col_idx, value=col_name)
        wb.save(self.excel_file)
        wb.close()

    def _ensure_sheet(self, sheet_name, headers):
        try:
            wb = openpyxl.load_workbook(self.excel_file)
        except Exception:
            return
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(sheet_name)
            ws.append(headers)
            wb.save(self.excel_file)
        wb.close()

    # ── Payment refs / Suspense sheet helpers ─────────────────────────────

    def _ensure_payment_refs_sheet_exists(self):
        try:
            wb = openpyxl.load_workbook(self.excel_file)
            if self.payment_refs_sheet not in wb.sheetnames:
                ws = wb.create_sheet(self.payment_refs_sheet)
                ws.append(self.PAYMENT_REF_HEADERS)
                wb.save(self.excel_file)
            wb.close()
        except Exception:
            pass

    def _ensure_suspense_sheet_exists(self):
        try:
            wb = openpyxl.load_workbook(self.excel_file)
            if self.suspense_sheet not in wb.sheetnames:
                ws = wb.create_sheet(self.suspense_sheet)
                ws.append(self.SUSPENSE_HEADERS)
                wb.save(self.excel_file)
            wb.close()
        except Exception:
            pass

    def _seed_default_categories(self):
        defaults = [
            "Maintenance", "Repairs", "Utilities", "Salaries", "Other",
            "Repair & Maintenance", "Service Charges", "Sinking Fund", "Interest",
        ]
        try:
            existing = self.get_expense_categories()
            existing_names = {c["Name"] for c in existing}
            for name in defaults:
                if name not in existing_names:
                    self.add_expense_category(name)
        except Exception:
            pass

    def _ensure_workbook_exists(self):
        if not self.excel_file.exists():
            wb = openpyxl.Workbook()
            wb.remove(wb.active)

            ws_members = wb.create_sheet(self.members_sheet)
            ws_members.append([
                "ID", "Plot_No", "Plot_Owner_Name", "Email",
                "Phone", "Pending_Interest",
            ])

            ws_ledger = wb.create_sheet(self.ledger_sheet)
            ws_ledger.append(self.LEDGER_HEADERS)

            ws_settings = wb.create_sheet(self.settings_sheet)
            ws_settings.append(["Key", "Value"])
            for row in [
                ["Current_FY", "2026-27"],
                ["Repair_Fund_Rate", "100.00"],
                ["Service_Charges_Rate", "885.00"],
                ["Sinking_Fund_Rate", "15.00"],
                ["April_1st_Auto_Entry_Amount", "12000"],
                ["Last_Voucher_No", "0"],
            ]:
                ws_settings.append(row)

            ws_refs = wb.create_sheet(self.payment_refs_sheet)
            ws_refs.append(self.PAYMENT_REF_HEADERS)

            ws_accounts = wb.create_sheet(self.accounts_sheet)
            ws_accounts.append(self.ACCOUNTS_HEADERS)

            wb.save(self.excel_file)

    # ── Members ────────────────────────────────────────────────────────────

    def get_all_members(self) -> List[Dict[str, Any]]:
        self._load_cache()
        return self._cache["members_list"]

    def get_member(self, member_id: int) -> Optional[Dict[str, Any]]:
        self._load_cache()
        return self._cache["members_by_id"].get(member_id)

    def get_member_by_plot_no(self, plot_no: str) -> Optional[Dict[str, Any]]:
        df = pd.read_excel(self.excel_file, sheet_name=self.members_sheet)
        member = df[df["Plot_No"].astype(str) == str(plot_no).zfill(2)]
        if member.empty:
            return None
        return member.iloc[0].to_dict()

    def add_member(self, member_data: Dict[str, Any]) -> int:
        df = pd.read_excel(self.excel_file, sheet_name=self.members_sheet)
        max_id = df["ID"].max() if len(df) > 0 else 0
        new_id = int(max_id) + 1
        member_data["ID"] = new_id
        new_row = pd.DataFrame([member_data])
        df = pd.concat([df, new_row], ignore_index=True)
        self._write_sheet(df, self.members_sheet)
        self._invalidate_cache()
        return new_id

    def update_member(self, member_id: int, member_data: Dict[str, Any]) -> bool:
        df = pd.read_excel(self.excel_file, sheet_name=self.members_sheet)
        mask = df["ID"] == member_id
        if not mask.any():
            return False
        for key, value in member_data.items():
            df.loc[mask, key] = value
        self._write_sheet(df, self.members_sheet)
        self._invalidate_cache()
        return True

    def _get_plot_no(self, member_id: int) -> str:
        member = self.get_member(member_id)
        return str(member.get("Plot_No", member_id)) if member else str(member_id)

    # ── Ledger ─────────────────────────────────────────────────────────────

    def get_member_ledger(self, member_id: int) -> List[Dict[str, Any]]:
        self._load_cache()
        cached = self._cache["ledger_by_mid"].get(member_id)
        if cached is not None:
            return cached
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.ledger_sheet)
        except Exception:
            return []
        if df.empty:
            return []
        if "Plot_No" in df.columns:
            member_df = df[df["Plot_No"] == member_id]
        else:
            member_df = df[df["Member_ID"] == member_id]
        return member_df.to_dict("records")

    def add_ledger_entry(self, member_id: int, entry: Dict[str, Any]) -> bool:
        entry["Member_ID"] = member_id
        entry["Plot_No"] = member_id
        entry_df = pd.DataFrame([entry])
        df = pd.read_excel(self.excel_file, sheet_name=self.ledger_sheet)
        df = _ensure_columns(df, entry_df.columns)
        df = pd.concat([df, entry_df.reindex(columns=df.columns)], ignore_index=True)
        self._write_sheet(df, self.ledger_sheet)
        self._invalidate_cache()
        return True

    def add_ledger_entries(self, entries: List[tuple]) -> int:
        df = pd.read_excel(self.excel_file, sheet_name=self.ledger_sheet)
        for member_id, entry in entries:
            entry["Member_ID"] = member_id
            entry["Plot_No"] = member_id
            entry_df = pd.DataFrame([entry])
            df = _ensure_columns(df, entry_df.columns)
            df = pd.concat([df, entry_df.reindex(columns=df.columns)], ignore_index=True)
        self._write_sheet(df, self.ledger_sheet)
        self._invalidate_cache()
        return len(entries)

    def get_ledger_entries(
        self, member_id: int, start_date: Optional[datetime] = None, end_date: Optional[datetime] = None
    ) -> List[Dict[str, Any]]:
        df = pd.read_excel(self.excel_file, sheet_name=self.ledger_sheet)
        df = df[df["Member_ID"] == member_id]
        if start_date or end_date:
            df["Date"] = pd.to_datetime(df["Date"], format="%d-%m-%Y", errors="coerce")
            if start_date:
                df = df[df["Date"] >= start_date]
            if end_date:
                df = df[df["Date"] <= end_date]
        return df.to_dict("records")

    def update_ledger_entry(self, member_id: int, vch_no: int, updates: Dict[str, Any]) -> bool:
        df = pd.read_excel(self.excel_file, sheet_name=self.ledger_sheet)
        mask = df["Vch_No"] == vch_no
        if not mask.any():
            return False
        for key, value in updates.items():
            df.loc[mask, key] = value
        self._write_sheet(df, self.ledger_sheet)
        self._invalidate_cache()
        return True

    def delete_ledger_entry(self, member_id: int, vch_no: int) -> bool:
        df = pd.read_excel(self.excel_file, sheet_name=self.ledger_sheet)
        mask = df["Vch_No"] == vch_no
        if not mask.any():
            return False
        df = df[~mask].reset_index(drop=True)
        self._write_sheet(df, self.ledger_sheet)
        self._invalidate_cache()
        return True

    # ── Settings ───────────────────────────────────────────────────────────

    def get_settings(self) -> Dict[str, Any]:
        df = pd.read_excel(self.excel_file, sheet_name=self.settings_sheet)
        settings = {}
        for _, row in df.iterrows():
            key = row["Key"]
            value = row["Value"]
            if isinstance(value, str):
                if value.isdigit():
                    settings[key] = int(value)
                elif value.replace(".", "", 1).isdigit():
                    settings[key] = float(value)
                else:
                    settings[key] = value
            else:
                settings[key] = value
        return settings

    def update_settings(self, settings: Dict[str, Any]) -> bool:
        df = pd.read_excel(self.excel_file, sheet_name=self.settings_sheet)
        for key, value in settings.items():
            mask = df["Key"] == key
            if mask.any():
                df.loc[mask, "Value"] = value
            else:
                new_row = pd.DataFrame([{"Key": key, "Value": value}])
                df = pd.concat([df, new_row], ignore_index=True)
        self._write_sheet(df, self.settings_sheet)
        self._invalidate_cache()
        return True

    def get_next_voucher_number(self) -> int:
        settings = self.get_settings()
        last_vch = int(settings.get("Last_Voucher_No", 0))
        next_vch = last_vch + 1
        self.update_settings({"Last_Voucher_No": str(next_vch)})
        return next_vch

    def get_next_voucher_numbers(self, count: int) -> List[int]:
        settings = self.get_settings()
        last_vch = int(settings.get("Last_Voucher_No", 0))
        numbers = list(range(last_vch + 1, last_vch + count + 1))
        self.update_settings({"Last_Voucher_No": str(numbers[-1])})
        return numbers

    # ── Outstanding ────────────────────────────────────────────────────────

    @staticmethod
    def _safe_float(val) -> float:
        if val is None:
            return 0.0
        try:
            v = float(val)
            return 0.0 if v != v else v
        except (ValueError, TypeError):
            return 0.0

    def get_current_outstanding(self, member_id: int) -> float:
        ledger = self.get_member_ledger(member_id)
        total_credit = sum(self._safe_float(e.get("Credit")) for e in ledger)
        total_debit = sum(self._safe_float(e.get("Debit")) for e in ledger)
        return total_credit - total_debit



    def get_member_ledger_all(self) -> List[Dict[str, Any]]:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.ledger_sheet)
        except Exception:
            return []
        if df.empty:
            return []
        return df.to_dict("records")

    def clear_auto_ledger_entries(self) -> int:
        """Delete ledger entries created by auto-processing (non-manual, non-invoice).
        Returns count of deleted entries. Preserves entries with 'Manual' in Description
        or Transaction_Type 'INVOICE'."""
        df = pd.read_excel(self.excel_file, sheet_name=self.ledger_sheet)
        if df.empty:
            return 0
        before = len(df)
        mask_manual = df["Description"].astype(str).str.contains("Manual", na=False)
        mask_invoice = df["Transaction_Type"].astype(str).str.upper().str.contains("INVOICE", na=False)
        df = df[mask_manual | mask_invoice].reset_index(drop=True)
        self._write_sheet(df, self.ledger_sheet)
        self._invalidate_cache()
        return before - len(df)

    def get_processed_statements(self) -> List[Dict[str, Any]]:
        try:
            df = pd.read_excel(self.excel_file, sheet_name="Processed_Statements")
        except Exception:
            return []
        if df.empty:
            return []
        return df.to_dict("records")

    def record_processed_statement(self, filename: str, date_range: str,
                                    entry_count: int, fmt: str, file_hash: str):
        try:
            df = pd.read_excel(self.excel_file, sheet_name="Processed_Statements")
        except Exception:
            df = pd.DataFrame(columns=["Filename", "Date_Range", "Processed_At",
                                        "Entry_Count", "Format", "File_Hash"])
        import datetime
        new_row = {
            "Filename": filename,
            "Date_Range": date_range,
            "Processed_At": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "Entry_Count": entry_count,
            "Format": fmt,
            "File_Hash": file_hash,
        }
        mask = (df["Filename"] == filename) & (df["File_Hash"] == file_hash)
        if mask.any():
            idx = mask.idxmax()
            for col, val in new_row.items():
                df.at[idx, col] = val
        else:
            df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
        self._write_sheet(df, "Processed_Statements")

    def refresh_reports_sheet(self):
        members = self.get_all_members()
        rows = []
        import datetime
        for m in members:
            mid = int(m.get("Member_ID", m.get("Plot_No", m.get("ID", 0))))
            outstanding = self.get_current_outstanding(mid)
            rows.append({
                "Plot_No": mid,
                "Owner_Name": m.get("Plot_Owner_Name", ""),
                "Current_Outstanding": outstanding,
                "Generated_At": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            })
        df = pd.DataFrame(rows)
        self._write_sheet(df, "Reports")

    # ── Payment References ─────────────────────────────────────────────────

    def record_payment_reference(self, member_id: int, identifier_type: str, identifier_value: str, date: str = None, transaction_type: str = None):
        if not identifier_value or not identifier_type:
            return
        val = str(identifier_value).strip()
        if not val:
            return
        self._ensure_payment_refs_sheet_exists()
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.payment_refs_sheet)
        except Exception:
            df = pd.DataFrame(columns=self.PAYMENT_REF_HEADERS)

        existing = df[(df["Member_ID"] == member_id) & (df["Identifier_Type"] == identifier_type) & (df["Identifier_Value"] == val)]
        if not existing.empty:
            df.loc[existing.index[0], "Last_Seen_Date"] = date or datetime.now().strftime("%d-%m-%Y")
            c = int(existing.iloc[0].get("Confidence", 0) or 0) + 1
            df.loc[existing.index[0], "Confidence"] = min(c, 100)
            if transaction_type:
                df.loc[existing.index[0], "Transaction_Type"] = transaction_type
        else:
            new_row = pd.DataFrame([{
                "Member_ID": member_id,
                "Identifier_Type": identifier_type,
                "Identifier_Value": val,
                "Last_Seen_Date": date or datetime.now().strftime("%d-%m-%Y"),
                "Confidence": 1,
                "Transaction_Type": transaction_type or "",
            }])
            df = pd.concat([df, new_row], ignore_index=True)

        self._write_sheet(df, self.payment_refs_sheet)

    def _migrate_payment_refs_schema(self):
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.payment_refs_sheet)
        except Exception:
            return
        changed = False
        if "Transaction_Type" not in df.columns:
            df["Transaction_Type"] = ""
            changed = True
        if not changed:
            return
        self._write_sheet(df, self.payment_refs_sheet)

    def _clean_payment_refs(self):
        import re
        self._migrate_payment_refs_schema()
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.payment_refs_sheet)
        except Exception:
            return {}
        if df.empty:
            return {"total_before": 0, "removed": 0, "details": {}}
        before = len(df)
        removed = {"MOBILE": 0, "PAYEE_NAME": 0, "UPI_ID": 0}
        mask = pd.Series(True, index=df.index)

        STOP_WORDS = {"MR", "MRS", "MS", "SHRI", "SMT", "DR", "SRI", "M/S", "CMN",
                      "PAY", "AND", "THE", "FROM", "B", "PAYMENT", "RECEIVED",
                      "TRANSFER", "NEFT", "RTGS", "IMPS", "CHQ", "CHEQUE",
                      "BANK", "REFERENCE", "CREDIT", "DEBIT", "BY", "TO"}

        for idx, row in df.iterrows():
            id_type = str(row.get("Identifier_Type", "") or "")
            val = str(row.get("Identifier_Value", "") or "").strip()

            if id_type == "MOBILE":
                if len(val) != 10 or not re.match(r'[6-9]\d{9}', val):
                    mask.at[idx] = False
                    removed["MOBILE"] += 1

            elif id_type == "PAYEE_NAME":
                is_amount = bool(re.match(r'^[\d,]+\.\d{2}$', val))
                is_ref_code = 5 <= len(val) <= 15 and bool(re.match(r'^[A-Z][A-Za-z0-9]*\d{4,}$', val))
                is_short_upper = val.isalpha() and val.isupper() and len(val) <= 4
                is_stop = val.upper() in STOP_WORDS
                if is_amount or is_ref_code or is_short_upper or is_stop:
                    mask.at[idx] = False
                    removed["PAYEE_NAME"] += 1

        df_clean = df[mask].reset_index(drop=True)
        removed_count = before - len(df_clean)
        if removed_count > 0:
            self._write_sheet(df_clean, self.payment_refs_sheet)
        return {
            "total_before": before,
            "removed": removed_count,
            "details": removed,
        }

    def find_member_by_identifier(self, identifier_value: str) -> Optional[Dict[str, Any]]:
        if not identifier_value:
            return None
        val = str(identifier_value).strip().upper()
        self._ensure_payment_refs_sheet_exists()
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.payment_refs_sheet)
        except Exception:
            return None
        if df.empty:
            return None
        df["match"] = df["Identifier_Value"].astype(str).str.strip().str.upper() == val
        matches = df[df["match"]]
        if matches.empty:
            return None
        best = matches.sort_values("Confidence", ascending=False).iloc[0]
        return self.get_member(int(best["Member_ID"]))

    def get_known_identifiers(self, member_id: int, min_confidence: int = 1) -> List[Dict[str, str]]:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.payment_refs_sheet)
        except Exception:
            return []
        if df.empty:
            return []
        member_refs = df[(df["Member_ID"] == member_id) & (df["Confidence"] >= min_confidence)]
        return member_refs.to_dict("records")

    def get_all_identifiers(self, id_type: Optional[str] = None) -> List[Dict[str, Any]]:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.payment_refs_sheet)
        except Exception:
            return []
        if df.empty:
            return []
        if id_type:
            df = df[df["Identifier_Type"] == id_type]
        return df.to_dict("records")

    def delete_identifier(self, ref_id: int) -> bool:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.payment_refs_sheet)
        except Exception:
            return False
        if df.empty:
            return False
        before = len(df)
        df = df.drop(index=ref_id).reset_index(drop=True) if ref_id < len(df) else df
        if len(df) == before:
            return False
        self._write_sheet(df, self.payment_refs_sheet)
        return True

    def add_identifier(self, member_id: int, id_type: str, id_value: str,
                       date: str = None, confidence: int = 1) -> int:
        self.record_payment_reference(member_id, id_type, id_value, date)
        return 0

    # ── Suspense Entries ──────────────────────────────────────────────────

    def add_suspense_entry(self, entry: Dict[str, Any]) -> int:
        self._ensure_suspense_sheet_exists()
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.suspense_sheet)
        except Exception:
            df = pd.DataFrame(columns=self.SUSPENSE_HEADERS)
        new_id = int(df["ID"].max()) + 1 if not df.empty and "ID" in df.columns else 1
        entry["ID"] = new_id
        entry["Entry_Date"] = entry.get("Entry_Date", datetime.now().strftime("%d-%m-%Y %H:%M"))
        entry["Status"] = entry.get("Status", "Pending")
        entry["Tagged_To"] = entry.get("Tagged_To", None)
        entry["Tagged_Date"] = entry.get("Tagged_Date", None)
        row = pd.DataFrame([{h: entry.get(h) for h in self.SUSPENSE_HEADERS}])
        if df.empty:
            df = row
        else:
            df = pd.concat([df, row], ignore_index=True)
        self._write_sheet(df, self.suspense_sheet)
        return new_id

    def get_suspense_entries(self, status: Optional[str] = None) -> List[Dict[str, Any]]:
        self._ensure_suspense_sheet_exists()
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.suspense_sheet)
        except Exception:
            return []
        if df.empty:
            return []
        if status:
            df = df[df["Status"] == status]
        return df.to_dict("records")

    def tag_suspense_entry(self, entry_id: int, member_id: int) -> bool:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.suspense_sheet)
        except Exception:
            return False
        if df.empty or "ID" not in df.columns:
            return False
        idx = df[df["ID"] == entry_id].index
        if idx.empty:
            return False
        for col in ["Status", "Tagged_To", "Tagged_Date"]:
            if col in df.columns and df[col].dtype.kind in ("i", "f"):
                df[col] = df[col].astype(object)
        df.loc[idx, "Status"] = "Tagged"
        df.loc[idx, "Tagged_To"] = member_id
        df.loc[idx, "Tagged_Date"] = datetime.now().strftime("%d-%m-%Y %H:%M")
        self._write_sheet(df, self.suspense_sheet)
        return True

    def delete_suspense_entry(self, entry_id: int) -> bool:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.suspense_sheet)
        except Exception:
            return False
        if df.empty or "ID" not in df.columns:
            return False
        before = len(df)
        df = df[df["ID"] != entry_id]
        if len(df) == before:
            return False
        self._write_sheet(df.reset_index(drop=True), self.suspense_sheet)
        return True

    # ── Expenses ───────────────────────────────────────────────────────────

    def add_expense(self, entry: Dict[str, Any]) -> int:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.expenses_sheet)
        except Exception:
            df = pd.DataFrame(columns=self.EXPENSES_HEADERS)
        new_id = int(df["ID"].max()) + 1 if not df.empty and "ID" in df.columns else 1
        entry["ID"] = new_id
        entry["Entry_Date"] = entry.get("Entry_Date", datetime.now().strftime("%d-%m-%Y %H:%M"))
        row = pd.DataFrame([{h: entry.get(h, "") for h in self.EXPENSES_HEADERS}])
        if df.empty:
            df = row
        else:
            df = pd.concat([df, row], ignore_index=True)
        self._write_sheet(df, self.expenses_sheet)
        return new_id

    def get_expenses(self, member_id: Optional[int] = None, from_date: Optional[str] = None,
                     to_date: Optional[str] = None, category: Optional[str] = None,
                     search: Optional[str] = None) -> List[Dict[str, Any]]:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.expenses_sheet)
        except Exception:
            return []
        if df.empty:
            return []
        if member_id is not None:
            df = df[df["Member_ID"] == member_id]
        if from_date:
            df = df[df["Date"] >= from_date]
        if to_date:
            df = df[df["Date"] <= to_date]
        if category:
            df = df[df["Category"] == category]
        if search:
            df = df[df.apply(lambda r: search.lower() in str(r.get("Particulars", "")).lower(), axis=1)]
        return df.to_dict("records")

    def delete_expense(self, expense_id: int) -> bool:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.expenses_sheet)
        except Exception:
            return False
        if df.empty or "ID" not in df.columns:
            return False
        before = len(df)
        df = df[df["ID"] != expense_id]
        if len(df) == before:
            return False
        self._write_sheet(df.reset_index(drop=True), self.expenses_sheet)
        return True

    def get_expense_summary(self, fy: str) -> Dict[str, float]:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.expenses_sheet)
        except Exception:
            return {}
        if df.empty:
            return {}
        summary = df.groupby("Category")["Amount"].sum().to_dict()
        return {str(k): float(v) for k, v in summary.items()}

    # ── Expense Categories ────────────────────────────────────────────────

    def get_expense_categories(self, parent: Optional[str] = None) -> List[Dict[str, Any]]:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.expense_categories_sheet)
        except Exception:
            return []
        if df.empty:
            return []
        if parent:
            df = df[df["Parent"] == parent]
        return df.to_dict("records")

    def add_expense_category(self, name: str, parent: Optional[str] = None) -> int:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.expense_categories_sheet)
        except Exception:
            df = pd.DataFrame(columns=self.EXPENSE_CATEGORIES_HEADERS)
        new_id = int(df["ID"].max()) + 1 if not df.empty and "ID" in df.columns else 1
        row = pd.DataFrame([{"ID": new_id, "Name": name, "Parent": parent}])
        if df.empty:
            df = row
        else:
            df = pd.concat([df, row], ignore_index=True)
        self._write_sheet(df, self.expense_categories_sheet)
        return new_id

    def delete_expense_category(self, category_id: int) -> bool:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.expense_categories_sheet)
        except Exception:
            return False
        if df.empty or "ID" not in df.columns:
            return False
        before = len(df)
        df = df[df["ID"] != category_id]
        if len(df) == before:
            return False
        self._write_sheet(df.reset_index(drop=True), self.expense_categories_sheet)
        return True

    # ── Interest Income ─────────────────────────────────────────────────

    def add_interest_income(self, entry: Dict[str, Any]) -> int:
        from datetime import datetime as _dt
        self._ensure_sheet(self.interest_income_sheet, self.INTEREST_INCOME_HEADERS)
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.interest_income_sheet)
        except Exception:
            df = pd.DataFrame(columns=self.INTEREST_INCOME_HEADERS)
        new_id = int(df["ID"].max()) + 1 if not df.empty and "ID" in df.columns else 1
        entry["ID"] = new_id
        entry["Entry_Date"] = entry.get("Entry_Date", _dt.now().strftime("%d-%m-%Y %H:%M"))
        row = pd.DataFrame([{h: entry.get(h, "") for h in self.INTEREST_INCOME_HEADERS}])
        if df.empty:
            df = row
        else:
            df = pd.concat([df, row], ignore_index=True)
        self._write_sheet(df, self.interest_income_sheet)
        return new_id

    def get_interest_income(self, source_file: Optional[str] = None) -> List[Dict[str, Any]]:
        self._ensure_sheet(self.interest_income_sheet, self.INTEREST_INCOME_HEADERS)
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.interest_income_sheet)
        except Exception:
            return []
        if df.empty:
            return []
        if source_file:
            df = df[df["Source_File"] == source_file]
        return df.to_dict("records")

    # ── Split Rules ────────────────────────────────────────────────────────

    def get_split_rules(self) -> List[Dict[str, Any]]:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.split_rules_sheet)
        except Exception:
            return []
        if df.empty:
            return []
        return df.to_dict("records")

    def add_split_rule(self, source_member_id: int, member_ids: List[int]) -> int:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.split_rules_sheet)
        except Exception:
            df = pd.DataFrame(columns=self.SPLIT_RULES_HEADERS)
        new_id = int(df["ID"].max()) + 1 if not df.empty and "ID" in df.columns else 1
        member_ids_str = ",".join(str(m) for m in member_ids)
        row = pd.DataFrame([{
            "ID": new_id,
            "Source_Member_ID": source_member_id,
            "Members": member_ids_str,
        }])
        if df.empty:
            df = row
        else:
            df = pd.concat([df, row], ignore_index=True)
        self._write_sheet(df, self.split_rules_sheet)
        return new_id

    def delete_split_rule(self, rule_id: int) -> bool:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.split_rules_sheet)
        except Exception:
            return False
        if df.empty or "ID" not in df.columns:
            return False
        before = len(df)
        df = df[df["ID"] != rule_id]
        if len(df) == before:
            return False
        self._write_sheet(df.reset_index(drop=True), self.split_rules_sheet)
        return True

    def get_split_group_for_member(self, member_id: int) -> Optional[List[int]]:
        rules = self.get_split_rules()
        for r in rules:
            if int(r["Source_Member_ID"]) == member_id:
                ids_str = str(r.get("Members", ""))
                if ids_str:
                    return [int(x.strip()) for x in ids_str.split(",") if x.strip()]
        return None

    # ── Ledger duplicate check ────────────────────────────────────────────

    def check_duplicate_ledger_entry(
        self, txn_id: str, date: str, amount: float, particulars: str = ""
    ) -> Optional[Dict[str, Any]]:
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.ledger_sheet)
        except Exception:
            return None
        if df.empty:
            return None
        txn_clean = str(txn_id or "").strip()
        raw_part = str(particulars or "").strip()
        part_clean = raw_part[3:] if raw_part.upper().startswith("BY ") else raw_part
        part_clean = part_clean[:40].upper()
        for _, e in df.iterrows():
            e_part_upper = str(e.get("Particulars", "") or "").upper()
            if "SPLIT FROM" in e_part_upper:
                continue
            e_date = str(e.get("Date", "") or "").strip()
            e_credit = float(e.get("Credit", 0) or 0)
            if e_date != date or abs(e_credit - amount) > 0.01:
                continue
            e_txn = str(e.get("Transaction_ID", "") or "").strip()
            if txn_clean and e_txn == txn_clean:
                return e.to_dict()
            e_part = str(e.get("Particulars", "") or "").strip()
            e_part_clean = e_part[3:] if e_part.upper().startswith("BY ") else e_part
            e_part_clean = e_part_clean[:40].upper()
            if part_clean and e_part_clean and (part_clean in e_part_clean or e_part_clean in part_clean):
                return e.to_dict()
        return None

    # ── Accounts Sheet ─────────────────────────────────────────────────

    def add_accounts_entries(self, entries: List[Dict[str, Any]]) -> int:
        df = pd.read_excel(self.excel_file, sheet_name=self.accounts_sheet)
        max_id = int(df["Entry_ID"].max()) + 1 if not df.empty and "Entry_ID" in df.columns else 1
        new_rows = []
        for i, entry in enumerate(entries):
            entry["Entry_ID"] = max_id + i
            entry.setdefault("Calculated_Balance", None)
            entry.setdefault("Balance_Check", "")
            entry.setdefault("Status", "Pending")
            new_rows.append({h: entry.get(h) for h in self.ACCOUNTS_HEADERS})
        new_df = pd.DataFrame(new_rows)
        if df.empty:
            df = new_df
        else:
            df = pd.concat([df, new_df], ignore_index=True)
        self._write_sheet(df, self.accounts_sheet)
        return len(entries)

    def get_accounts_entries(
        self, source_file: Optional[str] = None, status: Optional[str] = None
    ) -> List[Dict[str, Any]]:
        self._ensure_sheet(self.accounts_sheet, self.ACCOUNTS_HEADERS)
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.accounts_sheet)
        except Exception:
            return []
        if df.empty:
            return []
        if source_file:
            df = df[df["Source_File"] == source_file]
        if status:
            df = df[df["Status"] == status]
        return df.to_dict("records")

    def update_accounts_entry(self, entry_id: int, updates: Dict[str, Any]) -> bool:
        self._ensure_sheet(self.accounts_sheet, self.ACCOUNTS_HEADERS)
        df = pd.read_excel(self.excel_file, sheet_name=self.accounts_sheet)
        if df.empty or "Entry_ID" not in df.columns:
            return False
        idx = df[df["Entry_ID"] == entry_id].index
        if idx.empty:
            return False
        for key, value in updates.items():
            if key in df.columns:
                df.loc[idx, key] = value
        self._write_sheet(df, self.accounts_sheet)
        return True

    def delete_accounts_entries(self, entry_ids: List[int]) -> int:
        self._ensure_sheet(self.accounts_sheet, self.ACCOUNTS_HEADERS)
        df = pd.read_excel(self.excel_file, sheet_name=self.accounts_sheet)
        if df.empty or "Entry_ID" not in df.columns:
            return 0
        before = len(df)
        df = df[~df["Entry_ID"].isin(entry_ids)]
        deleted = before - len(df)
        if deleted:
            self._write_sheet(df.reset_index(drop=True), self.accounts_sheet)
        return deleted

    def get_accounts_summary(self) -> Dict[str, Any]:
        self._ensure_sheet(self.accounts_sheet, self.ACCOUNTS_HEADERS)
        try:
            df = pd.read_excel(self.excel_file, sheet_name=self.accounts_sheet)
        except Exception:
            return {"statements": [], "continuity": [], "gaps": [], "total_entries": 0}
        if df.empty:
            return {"statements": [], "continuity": [], "gaps": [], "total_entries": 0}
        result = {
            "total_entries": len(df),
            "by_status": df["Status"].value_counts().to_dict() if "Status" in df.columns else {},
            "balance_ok": int((df["Balance_Check"].astype(str).str.contains("✓")).sum()) if "Balance_Check" in df.columns else 0,
            "balance_mismatch": int((df["Balance_Check"].astype(str).str.contains("✗")).sum()) if "Balance_Check" in df.columns else 0,
            "statements": [],
            "continuity": [],
            "gaps": [],
        }
        if "Source_File" not in df.columns:
            return result
        file_dates = {}
        for sf in df["Source_File"].dropna().unique():
            sdf = df[df["Source_File"] == sf].sort_values("Statement_Seq")
            if sdf.empty:
                continue
            first_date = str(sdf.iloc[0].get("Date", ""))
            last_date = str(sdf.iloc[-1].get("Date", ""))
            opening = sdf.iloc[0].get("Balance")
            closing = sdf.iloc[-1].get("Balance")
            total_deposits = float(sdf["Deposit"].fillna(0).sum())
            total_withdrawals = float(sdf["Withdrawal"].fillna(0).sum())
            mismatches = int((sdf["Balance_Check"].astype(str).str.contains("✗")).sum())
            opening_val = float(opening) if opening is not None else 0
            closing_val = float(closing) if closing is not None else 0
            if opening_val != opening_val:
                opening_val = 0
            if closing_val != closing_val:
                closing_val = 0
            result["statements"].append({
                "source_file": sf,
                "first_date": first_date,
                "last_date": last_date,
                "entries": len(sdf),
                "opening_balance": opening_val,
                "closing_balance": closing_val,
                "total_deposits": total_deposits,
                "total_withdrawals": total_withdrawals,
                "mismatches": mismatches,
            })
            file_dates[sf] = {"first": first_date, "last": last_date}
        from datetime import datetime as _dt

        def _fy_sort_key(d_str):
            try:
                dt = _dt.strptime(str(d_str).strip(), "%d-%m-%Y")
            except (ValueError, TypeError):
                return (9999, "Unknown")
            if dt.month >= 4:
                return (dt.year, f"{str(dt.year)[2:]}-{str(dt.year+1)[2:]}")
            else:
                return (dt.year - 1, f"{str(dt.year-1)[2:]}-{str(dt.year)[2:]}")

        df2 = df.copy()
        df2["_fy_key"] = df2["Date"].apply(_fy_sort_key)
        df2["_fy"] = df2["_fy_key"].apply(lambda x: x[1])
        df2["_dt"] = pd.to_datetime(df2["Date"], format="%d-%m-%Y", errors="coerce")

        fy_groups = []
        for fy_key, fy_df in df2.groupby("_fy_key", sort=False):
            fy_df = fy_df.sort_values(["_dt", "Statement_Seq"])
            bounds_df = fy_df[fy_df["Source_File"] != "Legacy"]
            if bounds_df.empty:
                bounds_df = fy_df
            first = bounds_df.iloc[0]
            last = bounds_df.iloc[-1]
            first_dep = float(first.get("Deposit", 0) or 0)
            first_wd = float(first.get("Withdrawal", 0) or 0)
            if first_dep != first_dep:
                first_dep = 0.0
            if first_wd != first_wd:
                first_wd = 0.0
            bal_first = float(first.get("Balance", 0) or 0)
            if bal_first != bal_first:
                bal_first = 0.0
            opening_bal = bal_first - first_dep + first_wd
            closing_bal = float(last.get("Balance", 0) or 0)
            if opening_bal != opening_bal:
                opening_bal = 0
            if closing_bal != closing_bal:
                closing_bal = 0
            fy_groups.append({
                "fy": fy_key[1],
                "fy_start": fy_key[0],
                "opening_date": str(first.get("Date", "")),
                "closing_date": str(last.get("Date", "")),
                "opening_balance": opening_bal,
                "closing_balance": closing_bal,
                "entries": len(fy_df),
                "files": sorted(fy_df["Source_File"].dropna().unique().tolist()),
            })

        fy_groups.sort(key=lambda x: x["fy_start"])

        continuity = []
        gaps = []
        for i, fg in enumerate(fy_groups):
            rollover = None
            if i > 0:
                prev = fy_groups[i - 1]
                diff = round(prev["closing_balance"] - fg["opening_balance"], 2)
                is_match = abs(diff) < 0.01
                rollover = {
                    "from_fy": prev["fy"],
                    "from_closing_balance": prev["closing_balance"],
                    "from_closing_date": prev["closing_date"],
                    "to_opening_balance": fg["opening_balance"],
                    "to_opening_date": fg["opening_date"],
                    "match": is_match,
                    "difference": diff,
                }
                if not is_match:
                    gaps.append({
                        "from_fy": prev["fy"],
                        "to_fy": fg["fy"],
                        "expected_closing": prev["closing_balance"],
                        "actual_opening": fg["opening_balance"],
                        "difference": diff,
                    })
            continuity.append({
                "fy": fg["fy"],
                "opening_date": fg["opening_date"],
                "closing_date": fg["closing_date"],
                "opening_balance": fg["opening_balance"],
                "closing_balance": fg["closing_balance"],
                "entries": fg["entries"],
                "files": fg["files"],
                "rollover_from_prev": rollover,
            })

        result["continuity"] = continuity
        result["gaps"] = gaps
        return result

    def clear_accounts(self) -> int:
        df = pd.read_excel(self.excel_file, sheet_name=self.accounts_sheet)
        count = len(df)
        if count:
            self._write_sheet(df.iloc[:0], self.accounts_sheet)
        return count

    def migrate_accounts_from_ledger(self) -> int:
        if not self.accounts_sheet:
            return 0
        existing = self.get_accounts_entries()
        if existing:
            return 0
        entries = []
        ledger_all = self.get_member_ledger_all()
        seq = 0
        for e in ledger_all:
            date = str(e.get("Date", "") or "")
            credit = float(e.get("Credit", 0) or 0)
            debit = float(e.get("Debit", 0) or 0)
            if credit <= 0 and debit <= 0:
                continue
            seq += 1
            entries.append({
                "Date": date,
                "Particulars": str(e.get("Particulars", "") or ""),
                "Withdrawal": debit if debit > 0 else None,
                "Deposit": credit if credit > 0 else None,
                "Balance": None,
                "Transaction_Type": str(e.get("Transaction_Type", "") or ""),
                "Transaction_ID": str(e.get("Transaction_ID", "") or ""),
                "Status": "Matched",
                "Source_File": "Legacy",
                "Statement_Seq": seq,
            })
        suspense = self.get_suspense_entries()
        for s in suspense:
            seq += 1
            entries.append({
                "Date": str(s.get("Date", "") or ""),
                "Particulars": str(s.get("Particulars", "") or ""),
                "Withdrawal": None,
                "Deposit": float(s.get("Amount", 0) or 0),
                "Balance": None,
                "Transaction_Type": str(s.get("Transaction_Type", "") or ""),
                "Transaction_ID": str(s.get("Transaction_ID", "") or ""),
                "Status": "Unmatched",
                "Source_File": "Legacy",
                "Statement_Seq": seq,
            })
        expenses = self.get_expenses()
        for exp in expenses:
            seq += 1
            entries.append({
                "Date": str(exp.get("Date", "") or ""),
                "Particulars": str(exp.get("Particulars", "") or ""),
                "Withdrawal": float(exp.get("Amount", 0) or 0),
                "Deposit": None,
                "Balance": None,
                "Transaction_Type": str(exp.get("Transaction_Type", "") or ""),
                "Transaction_ID": str(exp.get("Transaction_ID", "") or ""),
                "Status": "Expensed",
                "Source_File": "Legacy",
                "Statement_Seq": seq,
            })
        if entries:
            self.add_accounts_entries(entries)
        return len(entries)

    # ── I/O ────────────────────────────────────────────────────────────────

    def _write_sheet(self, df: pd.DataFrame, sheet_name: str):
        wb = openpyxl.load_workbook(self.excel_file)
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)
        for col_idx, col_name in enumerate(df.columns, 1):
            ws.cell(row=1, column=col_idx, value=col_name)
        for row_idx in range(len(df)):
            for col_idx, col_name in enumerate(df.columns, 1):
                val = df.iloc[row_idx, col_idx - 1]
                if isinstance(val, float) and val != val:
                    val = None
                ws.cell(row=row_idx + 2, column=col_idx, value=val)
        wb.save(self.excel_file)
        wb.close()
